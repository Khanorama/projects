"""PathFinder backend (FastAPI)."""

import io, os, json, re
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
import openai
from dotenv import load_dotenv
from pypdf import PdfReader
import docx
from scraper import get_personalized_orgs

load_dotenv()

app = FastAPI(title="PathFinder API")
client = openai.OpenAI(api_key=os.getenv("OPENAI_API_KEY", ""))

AI_UNAVAILABLE = "AI service unavailable. Check OPENAI_API_KEY in backend/.env."


def ai_error(e: Exception) -> HTTPException:
    # 503 tells the frontend to use its fallback. Raising HTTPException (instead of
    # letting the error crash the request) also keeps the CORS headers on the response.
    print(f"OpenAI error: {e}")
    return HTTPException(status_code=503, detail=AI_UNAVAILABLE)


def unranked_result(orgs: list[dict]) -> dict:
    # used when the AI call fails: same shape as the AI result, scored by keyword relevance
    top = max((o.get("relevance_score", 0) for o in orgs), default=0) or 1
    return {
        "matches": [
            {
                "org_name":    o["name"],
                "type":        o["type"],
                "xp_category": o["xp_category"],
                "xp_value":    o["xp_value"],
                "why_it_fits": o["description"],
                "next_action": "Visit the website to see how to get involved.",
                "contact":     o["contact"],
                "website":     o["website"],
                "location":    o["location"],
                "tags":        o["tags"],
                "match_score": 60 + round(35 * o.get("relevance_score", 0) / top),
            }
            for o in orgs
        ],
        "coach_message": (
            "AI ranking is unavailable right now, so these are the raw search results, "
            "sorted by how closely they match your profile keywords."
        ),
        "ai_ranked": False,
    }

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


def extract_resume_text(filename: str, contents: bytes) -> str:
    """Text from a PDF or .docx resume. Raises an HTTPException the student can read if it fails."""
    name = (filename or "").lower()
    if not name.endswith((".pdf", ".docx")):
        raise HTTPException(415, "Please upload a PDF or Word (.docx) file. Older .doc files need to be saved as .docx first.")
    try:
        if name.endswith(".pdf"):
            parts = [page.extract_text() or "" for page in PdfReader(io.BytesIO(contents)).pages]
        else:
            doc = docx.Document(io.BytesIO(contents))
            parts = [para.text for para in doc.paragraphs]
            # a lot of resume templates put sections in tables
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    except Exception as e:
        print(f"Resume read error: {e}")
        raise HTTPException(422, "That file couldn't be opened. Check that it isn't damaged or password-protected.")
    text = "\n".join(part for part in parts if part.strip()).strip()
    if not text:
        raise HTTPException(422, "No text was found in that file. If it's a scanned PDF, upload a text-based PDF or a .docx instead.")
    return text


def parse_ai_json(raw: str) -> dict | None:
    # the model sometimes wraps the JSON in ``` or adds a sentence before it
    match = re.search(r"\{.*\}", raw or "", re.DOTALL)
    if not match:
        return None
    try:
        data = json.loads(match.group(0))
    except json.JSONDecodeError:
        return None
    return data if isinstance(data, dict) else None


@app.post("/parse-resume")
async def parse_resume(file: UploadFile = File(...)):
    resume_text = extract_resume_text(file.filename, await file.read())

    schema = (
        'Return ONLY valid JSON, no markdown:\n'
        '{"name":"","grade":"","school":"","location":"","skills":[],'
        '"interests":[],"experience":[],"career_goals":""}'
    )

    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": "Extract structured info from resumes. Return only valid JSON, no markdown."},
                {"role": "user",   "content": f"{schema}\n\nResume text:\n{resume_text[:6000]}"},
            ],
            temperature=0.1, max_tokens=500,
        )
    except openai.OpenAIError as e:
        raise ai_error(e)

    data = parse_ai_json(response.choices[0].message.content)
    if data is None:
        raise HTTPException(502, "The AI's reply couldn't be read. Please try uploading your resume again.")
    return {"resume_data": data}


@app.post("/search/stream")
async def search_stream(payload: dict):
    """Server-sent events: {"type": "thought" | "result" | "error", "payload": ...}"""
    student = payload.get("student", {})

    def event_stream():
        def send(t, c):
            return f"data: {json.dumps({'type': t, 'payload': c})}\n\n"

        name      = (student.get("name") or "you").split()[0]
        interests = ", ".join(student.get("interests") or []) or "your interests"
        goals     = student.get("career_goals") or ""
        location  = student.get("location") or "your area"

        yield send("thought", f"Initializing PathFinder for {name}...")
        yield send("thought", f"Interests detected: {interests}")
        yield send("thought", f"Searching for real local opportunities near {location}...")
        yield send("thought", "Running web searches tailored to your profile...")

        try:
            orgs = get_personalized_orgs(student)
            if orgs:
                yield send("thought", f"Found {len(orgs)} local opportunities.")
            else:
                yield send("thought", "The search found nothing, so asking the AI for suggestions instead...")
        except Exception as e:
            print(f"Scrape error: {e}")
            orgs = []
            yield send("thought", "The search failed, so asking the AI for suggestions instead...")

        if goals:
            yield send("thought", f"Matching against your goals: {goals}")
        yield send("thought", "Ranking all opportunities by interest and career fit...")

        # the prompt depends on whether the search found anything
        if orgs:
            data_section = (
                f"Scraped live opportunities (ALL must be included in results):\n"
                f"{json.dumps(orgs, indent=2)}\n\n"
                f"Return ALL {len(orgs)} opportunities. "
                f"Do NOT filter any out based on schedule or availability; "
                f"only exclude if the student stated an explicit hard restriction. "
                f"Copy each org's 'website', 'type' and 'location' fields exactly as-is into your JSON."
            )
        else:
            data_section = (
                f"No scraped data was available (web search returned no results).\n\n"
                f"Use your own knowledge to suggest 6-8 REAL volunteer and internship opportunities "
                f"that actually exist near {location} for this student. "
                f"Include real org names, real websites, real contact info where you know it. "
                f"Focus on opportunities relevant to: {interests}. "
                f"Set website to the real URL of each org."
            )

        try:
            stream = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "You are PathFinder, an encouraging AI career coach for high school students. "
                            "First think out loud in 3-4 short sentences referencing the student by "
                            "first name and mentioning specific org names. "
                            "Then write exactly '---RESULT---' on its own line, followed IMMEDIATELY "
                            "by valid JSON (no markdown fences):\n"
                            '{"matches":[' +
                            '{"org_name":"","type":"Internship or Volunteer","location":"",' +
                            '"xp_category":"","xp_value":0,' +
                            '"why_it_fits":"2 sentences personal to this student",' +
                            '"next_action":"one concrete step the student can take today",' +
                            '"contact":"",' +
                            '"website":"",' +
                            '"tags":[]}' +
                            '],' +
                            '"coach_message":"2-3 motivational sentences"}'
                        ),
                    },
                    {
                        "role": "user",
                        "content": (
                            f"Student profile:\n{json.dumps(student, indent=2)}\n\n"
                            + data_section +
                            f"\n\nBe specific about WHY each org fits this particular student's "
                            f"interests, skills, and career goals."
                        ),
                    },
                ],
                stream=True,
                temperature=0.5,
                max_tokens=3000,
            )

            buffer = ""
            result_mode = False

            for chunk in stream:
                delta = chunk.choices[0].delta.content or ""
                buffer += delta

                if "---RESULT---" in buffer and not result_mode:
                    thoughts_raw, buffer = buffer.split("---RESULT---", 1)
                    for sentence in re.split(r"(?<=[.!?])\s+", thoughts_raw.strip()):
                        if sentence.strip():
                            yield send("thought", sentence.strip())
                    result_mode = True
                    continue

                if not result_mode:
                    parts = re.split(r"(?<=[.!?])\s+", buffer)
                    if len(parts) > 1:
                        for sentence in parts[:-1]:
                            if sentence.strip():
                                yield send("thought", sentence.strip())
                        buffer = parts[-1]

            if result_mode and buffer.strip():
                clean = re.sub(r"^```(?:json)?\s*", "", buffer.strip())
                clean = re.sub(r"\s*```$", "", clean).strip()
                try:
                    result = json.loads(clean)
                    # fill in any website the AI left out from the scraped data
                    org_map = {o["name"].lower(): o for o in orgs}
                    for m in result.get("matches", []):
                        if not m.get("website"):
                            key = m.get("org_name", "").lower()
                            for k, v in org_map.items():
                                if key in k or k in key:
                                    m["website"] = v.get("website", "")
                                    break
                    n = len(result.get("matches", []))
                    yield send("thought", f"Done. Found {n} opportunities for {name}.")
                    yield send("result", result)
                except Exception:
                    raw_match = re.search(r"\{.*\}", clean, re.DOTALL)
                    if raw_match:
                        try:
                            yield send("result", json.loads(raw_match.group(0)))
                        except Exception as e:
                            yield send("error", f"Parse error: {e}")
                    else:
                        yield send("error", "Could not parse results. Please try again.")
            elif not result_mode:
                # the reply never had the ---RESULT--- marker
                yield send("error", "No structured result returned. Please try again.")
        except Exception as e:
            # usually a missing or invalid API key
            print(f"AI ranking error: {e}")
            if orgs:
                yield send("thought", "AI ranking is unavailable, so showing the raw search results.")
                yield send("result", unranked_result(orgs))
            else:
                yield send("error", f"Search failed: {AI_UNAVAILABLE}")

        yield "data: [DONE]\n\n"

    return StreamingResponse(event_stream(), media_type="text/event-stream")


@app.post("/draft-email")
async def draft_email(payload: dict):
    """Outreach email for one opportunity (the Draft email button and the Workshop email tool)."""
    student = payload.get("student", {})
    match   = payload.get("match", {})
    custom  = payload.get("custom_prompt", "")

    name      = student.get("name") or "Student"
    school    = student.get("school") or ""
    level     = "college" if "college" in str(student.get("grade") or "").lower() else "high school"
    interests = ", ".join(student.get("interests") or [])
    skills    = ", ".join(student.get("skills") or [])
    schedule  = student.get("schedule", "weekends")
    goals     = student.get("career_goals", "")
    org_name  = (match.get("org_name") or match.get("organization") or
                 match.get("title") or custom or "the organization")
    why_fits  = match.get("why_it_fits", "")
    contact   = match.get("contact", "")
    next_step = match.get("next_action", "")
    website   = match.get("website", "")

    prompt = (
        f"Write a short, genuine outreach email from a {level} student "
        f"to the volunteer or internship coordinator at {org_name}.\n\n"
        f"Student: {name}, {level} student" + (f" at {school}" if school else "") + "\n"
        f"Interests: {interests}\nSkills: {skills}\n"
        f"Availability: {schedule}\nCareer goals: {goals}\n\n"
        f"Why this opportunity fits the student: {why_fits}\n"
        f"Contact info available: {contact}\n"
        f"Organisation website: {website}\n"
        f"Suggested next step: {next_step}\n\n"
        f"Rules:\n"
        f"- First line must be 'Subject: ...'\n"
        f"- 3 short paragraphs, under 180 words total\n"
        f"- Professional but warm, not a template, not robotic\n"
        f"- Reference something specific about {org_name}\n"
        f"- End with a clear, low-friction ask (e.g. 'Would it be possible to...')\n"
        f"- Sign off with the student's name" + (" and school" if school else "") + "\n"
    )

    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are a career coach helping high school students write "
                        "outreach emails. Always start with 'Subject:' on the first line. "
                        "Make the email feel genuinely written by this specific student."
                    ),
                },
                {"role": "user", "content": prompt},
            ],
            temperature=0.55,
            max_tokens=400,
        )
    except openai.OpenAIError as e:
        raise ai_error(e)
    return {"email": response.choices[0].message.content}


@app.post("/workshop")
async def workshop(payload: dict):
    tool    = payload.get("tool", "guidance")
    prompt  = payload.get("prompt", "")
    student = payload.get("student", {})

    name      = student.get("name", "the student")
    skills    = ", ".join(student.get("skills") or []) or "various skills"
    interests = ", ".join(student.get("interests") or []) or "multiple interests"
    goals     = student.get("career_goals", "")

    if tool == "essay":
        system = "You are a college application coach. Be specific, warm, and avoid clichés."
        user   = (f"Student: {name}\nSkills: {skills}\nInterests: {interests}\nGoals: {goals}\n\n"
                  f"Essay request: {prompt}\n\nProvide: title, 2-sentence approach, opening direction, draft body paragraph.")
    else:
        system = "You are a career counsellor. Give a practical, encouraging action plan."
        user   = (f"Student: {name}\nSkills: {skills}\nInterests: {interests}\nGoals: {goals}\n\n"
                  f"Request: {prompt}\n\nProvide: title, 1-sentence summary, 'Recommended approach', 'Next three moves'.")

    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role": "system", "content": system}, {"role": "user", "content": user}],
            temperature=0.6, max_tokens=600,
        )
    except openai.OpenAIError as e:
        raise ai_error(e)

    raw   = response.choices[0].message.content
    lines = raw.strip().split("\n")
    title = lines[0].strip("#* ") if lines else "AI Response"
    summary = lines[1].strip() if len(lines) > 1 else ""

    sections, label, text = [], None, []
    for line in lines[2:]:
        s = line.strip()
        if (s.startswith("**") and s.endswith("**") and len(s) < 60) or s.startswith("##"):
            if label: sections.append({"label": label, "text": " ".join(text).strip()})
            label, text = s.strip("*#").strip(), []
        elif s:
            text.append(s)
    if label and text:
        sections.append({"label": label, "text": " ".join(text).strip()})
    if not sections:
        sections = [{"label": "Response", "text": raw}]

    def clean(text):
        text = re.sub(r'^\*\s+', '- ', text, flags=re.MULTILINE) # * bullets (before italics, or "* a *b*" breaks)
        text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)   # **bold** -> plain
        text = re.sub(r'\*([^*]+)\*',   r'\1', text)    # *italic* -> plain
        text = re.sub(r'^#+\s+', '', text, flags=re.MULTILINE)  # ## headings
        return text.strip()

    return {
        "title":    clean(title),
        "summary":  clean(summary or f"Generated for {name}."),
        "sections": [{"label": clean(s["label"]), "text": clean(s["text"])} for s in sections],
    }


@app.get("/health")
def health():
    return {"status": "ok"}
